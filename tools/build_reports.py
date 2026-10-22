"""Generuoja ataskaitų failus (Word, PDF, Excel) iš reports/<emitentas>/<laikotarpis>.yaml ir data/history/*.csv.

Kalbos ir formatai - config/settings.yaml. Rezultatas: docs/files/<emitentas>/<emitentas>_<laikotarpis>_<kalba>.<docx|pdf|xlsx>
Naudojimas: python tools/build_reports.py            - visos ataskaitos
            python tools/build_reports.py ignitis    - tik vienas emitentas
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common as C  # noqa: E402

FONT_DIR = C.ROOT / "assets" / "fonts"
pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
ACCENT = "1F4E79"
STATUS_KEYS = {"yes": "st_yes", "no": "st_no", "review": "st_review", "found": "st_found", "not_found": "st_not_found",
               "met": "st_met", "partly": "st_partly", "not_met": "st_not_met", "pending": "st_pending", "changed": "st_changed"}


# ---------- dokumento modelis (bendras Word ir PDF) ----------
def model(rep, lang):
    L = C.i18n()
    M = C.metrics()
    t = lambda k: C.tr(L.get(k, k), lang)  # noqa: E731
    iss = C.issuers()[rep["issuer"]]
    hist = C.history(rep["issuer"])
    secs = [s for s in C.securities() if s["issuer"] == rep["issuer"] and s.get("held") == "yes"]
    blocks = [("title", f"{iss['name']} – {t('title')}"),
              ("meta", [(t("analysis_date"), str(rep["analysis_date"])), (t("period"), rep["period"]),
                        ("Reitingas / Rating", rep.get("rating", "")), ("Kita ataskaita / Next report", rep.get("next_report", ""))])]

    # priemonės
    rows = []
    for s in secs:
        mat = s["maturity_date"] or (t("perpetual") if s["type"] == "bond" else t("share"))
        if s["status"] == "matured":
            mat += " (išpirkta)"
        cpn = (s["coupon_pct"] + " %") if s["coupon_pct"] else "–"
        note = s.get("status_note", "")
        if s.get("verified") == "review":
            note = ("⚠ " + note) if note else "⚠"
        rows.append([s["isin"], s["name"], cpn, mat, note])
    blocks += [("h2", t("instruments")), ("table", [t("isin"), t("instrument"), t("coupon"), t("maturity"), t("note")], rows, [3.0, 4.2, 1.6, 2.2, 6.0])]

    # rodikliai
    periods = rep.get("periods", [])
    rows = []
    for m in rep.get("ratios", []):
        vals = hist.get(m, {})
        if not vals:
            continue
        unit = next(iter(vals.values()))["unit"]
        label = C.tr(M.get(m, m), lang)
        ul = C.unit_label(unit, lang)
        if ul:
            label += f", {ul}"
        cells = []
        for p in periods:
            r = vals.get(p)
            if not r:
                cells.append("–")
                continue
            v = C.fmt_num(r["value"], unit, lang)
            if r.get("verified") == "check":
                v += "*"
            elif r.get("verified") == "calc":
                v += "ᶜ"
            cells.append(v)
        src = sorted({f"{r['source']}" + (f" p.{r['page']}" if r.get("page") else "") for r in vals.values() if r["source"] != "calc"})
        rows.append([label] + cells + ["; ".join(src[:2])])
    w = [5.0] + [2.1] * len(periods) + [2.6]
    blocks += [("h2", t("key_ratios")), ("table", [t("metric")] + periods + [t("sources")], rows, w),
               ("small", "* – reikia patikrinti originale / verify in source;  ᶜ – apskaičiuota įrankio / calculated by the tool")]

    for key in ("quick_summary", "business", "financials"):
        blocks += [("h2", t(key)), ("bullets", [C.tr(x, lang) for x in rep.get(key, [])])]

    rows = [[C.tr(p["plan"], lang), C.tr(p["outcome"], lang), t(STATUS_KEYS.get(p["status"], p["status"])), p.get("src", "")] for p in rep.get("plans", [])]
    blocks += [("h2", t("plans")), ("table", [t("plan"), t("outcome"), t("status"), t("sources")], rows, [6.0, 6.0, 2.4, 2.6])]

    blocks += [("h2", t("investor")), ("bullets", [C.tr(x, lang) for x in rep.get("investor", [])])]

    # ESG
    seb = C.seb()
    esg = rep.get("esg", {})
    pol = [seb["policies"][k]["name"] + ("" if seb["policies"][k].get("verified") == "yes" else " (turinys nepatikrintas / not verified)")
           for k in esg.get("policies", []) if k in seb["policies"]]
    blocks += [("h2", t("esg")), ("p", "SEB: " + "; ".join(pol))]
    items = esg.get("items", {})
    rows = []
    group = None
    for q in seb["checklist"]:
        if q["group"] == "manual":
            continue
        if q["group"] != group:
            group = q["group"]
            rows.append([("§", t("grp_" + group))])
        it = items.get(q["id"], {"status": "review", "evidence": ""})
        rows.append([C.tr(q["q"], lang), t(STATUS_KEYS.get(it["status"], it["status"])), C.tr(it.get("evidence"), lang)])
    blocks += [("table", [t("question"), t("status"), t("evidence")], rows, [5.0, 2.8, 9.2]),
               ("p", f"{t('esg_conclusion')}: {C.tr(esg.get('conclusion'), lang)}")]

    # rankinės patikros
    rows = []
    for k in ("sanctions", "media"):
        ch = rep.get("checks", {}).get(k, {})
        q = next(C.tr(x["q"], lang) for x in seb["checklist"] if x["id"] == k)
        rows.append([q, t(STATUS_KEYS.get(ch.get("status", "review"), "st_review")), C.tr(ch.get("note"), lang)])
    blocks += [("h2", t("checks")), ("table", [t("question"), t("status"), t("note")], rows, [5.0, 2.8, 9.2])]

    rows = [[k, v["title"], str(v.get("date", "")), v["url"]] for k, v in rep.get("sources", {}).items()]
    blocks += [("h2", t("sources")), ("table", ["#", "", "", "URL"], rows, [1.0, 5.5, 2.0, 8.5]), ("small", t("disclaimer"))]
    return blocks


# ---------- Word ----------
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def to_docx(blocks, path):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for kind, *rest in blocks:
        if kind == "title":
            h = doc.add_heading(rest[0], level=0)
            h.runs[0].font.color.rgb = RGBColor.from_string(ACCENT)
        elif kind == "h2":
            h = doc.add_heading(rest[0], level=2)
            h.runs[0].font.color.rgb = RGBColor.from_string(ACCENT)
        elif kind == "meta":
            for k, v in rest[0]:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.add_run(f"{k}: ").bold = True
                p.add_run(v)
        elif kind == "p":
            doc.add_paragraph(rest[0])
        elif kind == "small":
            p = doc.add_paragraph()
            r = p.add_run(rest[0])
            r.font.size = Pt(8)
            r.italic = True
        elif kind == "bullets":
            for b in rest[0]:
                p = doc.add_paragraph(b, style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
        elif kind == "table":
            header, rows, widths = rest
            tbl = doc.add_table(rows=1, cols=len(header))
            tbl.style = "Table Grid"
            for i, hcell in enumerate(header):
                c = tbl.rows[0].cells[i]
                c.text = ""
                run = c.paragraphs[0].add_run(str(hcell))
                run.bold = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
                shade(c, ACCENT)
            for row in rows:
                cells = tbl.add_row().cells
                if row and isinstance(row[0], tuple):  # grupės eilutė
                    m = cells[0].merge(cells[-1])
                    m.text = ""
                    m.paragraphs[0].add_run(row[0][1]).bold = True
                    shade(m, "DCE6F1")
                    continue
                for i, v in enumerate(row):
                    cells[i].text = ""
                    run = cells[i].paragraphs[0].add_run(str(v))
                    run.font.size = Pt(8.5)
                    if i > 0 and len(header) > 4 and header[0] != "#" and str(v)[:1].isdigit():
                        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for row in tbl.rows:
                for i, wcm in enumerate(widths[:len(row.cells)]):
                    row.cells[i].width = Cm(wcm)
            doc.add_paragraph()
    doc.save(path)


# ---------- PDF ----------
def to_pdf(blocks, path):
    S = lambda name, **kw: ParagraphStyle(name, fontName=kw.pop("font", "DejaVu"), **kw)  # noqa: E731
    st = {"title": S("t", font="DejaVu-Bold", fontSize=16, leading=20, textColor=colors.HexColor("#" + ACCENT), spaceAfter=8),
          "h2": S("h", font="DejaVu-Bold", fontSize=11.5, leading=15, textColor=colors.HexColor("#" + ACCENT), spaceBefore=10, spaceAfter=4),
          "p": S("p", fontSize=9, leading=12.5, alignment=TA_LEFT),
          "b": S("b", fontSize=9, leading=12.5, leftIndent=10, bulletIndent=2, spaceAfter=2),
          "small": S("s", fontSize=7, leading=9, textColor=colors.HexColor("#555555"), spaceBefore=3),
          "cell": S("c", fontSize=7.5, leading=9.5), "cellb": S("cb", font="DejaVu-Bold", fontSize=7.5, leading=9.5, textColor=colors.white),
          "cellg": S("cg", font="DejaVu-Bold", fontSize=7.5, leading=9.5)}
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    flow = []
    for kind, *rest in blocks:
        if kind == "title":
            flow.append(Paragraph(esc(rest[0]), st["title"]))
        elif kind == "h2":
            flow.append(Paragraph(esc(rest[0]), st["h2"]))
        elif kind == "meta":
            for k, v in rest[0]:
                flow.append(Paragraph(f"<b>{esc(k)}:</b> {esc(v)}", st["p"]))
        elif kind == "p":
            flow.append(Paragraph(esc(rest[0]), st["p"]))
        elif kind == "small":
            flow.append(Paragraph(esc(rest[0]), st["small"]))
        elif kind == "bullets":
            for b in rest[0]:
                flow.append(Paragraph(esc(b), st["b"], bulletText="•"))
        elif kind == "table":
            header, rows, widths = rest
            total = sum(widths)
            widths = [w * 17.4 / total * cm for w in widths]
            data = [[Paragraph(esc(h), st["cellb"]) for h in header]]
            style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + ACCENT)),
                     ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0B7C3")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
            for row in rows:
                if row and isinstance(row[0], tuple):
                    i = len(data)
                    data.append([Paragraph(esc(row[0][1]), st["cellg"])] + [""] * (len(header) - 1))
                    style += [("SPAN", (0, i), (-1, i)), ("BACKGROUND", (0, i), (-1, i), colors.HexColor("#DCE6F1"))]
                else:
                    data.append([Paragraph(esc(v), st["cell"]) for v in row])
            tb = Table(data, colWidths=widths, repeatRows=1)
            tb.setStyle(TableStyle(style))
            flow += [tb, Spacer(1, 4)]
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            title=blocks[0][1], author="issuers-review")
    doc.build(flow)


# ---------- Excel ----------
def to_xlsx(rep, lang, path):
    L = C.i18n()
    M = C.metrics()
    t = lambda k: C.tr(L.get(k, k), lang)  # noqa: E731
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=ACCENT)

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    # 1. rodikliai: skaičiai kaip skaičiai (galima skaičiuoti)
    hist = C.history(rep["issuer"])
    periods = sorted({p for v in hist.values() for p in v}, key=lambda p: (p[-4:], p))
    ws = wb.active
    ws.title = t("key_ratios")[:31]
    rows = []
    for m, vals in hist.items():
        unit = next(iter(vals.values()))["unit"]
        row = [C.tr(M.get(m, m), lang), unit]
        for p in periods:
            r = vals.get(p)
            row.append(float(r["value"]) if r and r["value"] else None)
        rows.append(row)
    sheet(ws, [t("metric"), "Unit"] + periods, rows, [34, 8] + [11] * len(periods))

    ws = wb.create_sheet(t("instruments")[:31])
    secs = [s for s in C.securities() if s["issuer"] == rep["issuer"]]
    sheet(ws, ["ISIN", t("instrument"), "Held", t("coupon"), "Issue", t("maturity"), "Call", t("status"), t("note"), "Source"],
          [[s["isin"], s["name"], s["held"], float(s["coupon_pct"]) if s["coupon_pct"] else None, s["issue_date"], s["maturity_date"],
            s["call_date"], s["status"], s["status_note"], s["source_url"]] for s in secs], [15, 34, 6, 8, 11, 11, 11, 9, 50, 40])

    ws = wb.create_sheet(t("plans")[:31])
    sheet(ws, [t("plan"), t("outcome"), t("status"), t("sources")],
          [[C.tr(p["plan"], lang), C.tr(p["outcome"], lang), t(STATUS_KEYS.get(p["status"], p["status"])), p.get("src", "")] for p in rep.get("plans", [])],
          [55, 55, 14, 16])

    ws = wb.create_sheet("ESG")
    seb = C.seb()
    items = rep.get("esg", {}).get("items", {})
    sheet(ws, [t("question"), t("status"), t("evidence")],
          [[C.tr(q["q"], lang), t(STATUS_KEYS.get(items.get(q["id"], {}).get("status", "review"), "st_review")), C.tr(items.get(q["id"], {}).get("evidence"), lang)]
           for q in seb["checklist"] if q["group"] != "manual"], [50, 18, 90])

    ws = wb.create_sheet(t("sources")[:31])
    sheet(ws, ["#", "Title", "Date", "URL"], [[k, v["title"], str(v.get("date", "")), v["url"]] for k, v in rep.get("sources", {}).items()], [6, 60, 12, 90])
    wb.save(path)


def build(only=None):
    cfg = C.settings()
    built = []
    for issuer in C.issuers():
        if only and issuer != only:
            continue
        for rep in C.reports(issuer):
            out = C.DOCS / "files" / issuer
            out.mkdir(parents=True, exist_ok=True)
            for lang in cfg.get("languages", ["lt"]):
                base = out / f"{issuer}_{rep['period']}_{lang}"
                blocks = model(rep, lang)
                if "docx" in cfg["formats"]:
                    to_docx(blocks, base.with_suffix(".docx"))
                if "pdf" in cfg["formats"]:
                    to_pdf(blocks, base.with_suffix(".pdf"))
                if "xlsx" in cfg["formats"]:
                    to_xlsx(rep, lang, base.with_suffix(".xlsx"))
                built.append(base.name)
    print("Sugeneruota:", ", ".join(built) or "nieko")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else None)
